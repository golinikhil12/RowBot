"""
exporter.py
Generates downloadable CSV and Excel files from a cleaned DataFrame.
Handles CSV formula injection mitigation for user-controlled text.
"""
import io
import re
import csv
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime, timezone


# Characters that spreadsheet apps interpret as formula starters
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def sanitize_cell(value) -> str:
    """Prefix formula-looking values with a single quote to prevent injection."""
    if isinstance(value, str) and value and value[0] in FORMULA_PREFIXES:
        return "'" + value
    return value


def export_csv(df: pd.DataFrame, original_filename: str) -> dict:
    """
    Export DataFrame to CSV bytes.
    Returns: {ok, bytes, filename, mimetype}
    """
    df_export = df.map(sanitize_cell)
    buf = io.StringIO()
    df_export.to_csv(buf, index=False, quoting=csv.QUOTE_NONNUMERIC)
    csv_bytes = buf.getvalue().encode("utf-8-sig")  # BOM for Excel compatibility

    stem = re.sub(r"\.[^.]+$", "", original_filename)
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    filename = f"{stem}_cleaned_{ts}.csv"

    return {
        "ok": True,
        "bytes": csv_bytes,
        "filename": filename,
        "mimetype": "text/csv",
    }


def export_excel(df: pd.DataFrame, original_filename: str) -> dict:
    """
    Export DataFrame to Excel bytes with basic styling.
    Returns: {ok, bytes, filename, mimetype}
    """
    buf = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cleaned Data"

    # Style constants
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    body_font = Font(name="Calibri", size=10)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(border_style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Write header
    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=str(col_name))
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = border

    ws.row_dimensions[1].height = 20

    # Write data rows
    for row_idx, row in enumerate(df.itertuples(index=False), start=2):
        for col_idx, value in enumerate(row, start=1):
            # Sanitize formula injection
            safe_val = sanitize_cell(value)
            cell = ws.cell(row=row_idx, column=col_idx, value=safe_val)
            cell.font = body_font
            cell.border = border
            # Alternate row fill
            if row_idx % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F5F8FC")

    # Auto-size columns
    for col_cells in ws.columns:
        max_len = 0
        col_letter = col_cells[0].column_letter
        for cell in col_cells:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 40)

    # Freeze header row
    ws.freeze_panes = ws["A2"]

    wb.save(buf)
    excel_bytes = buf.getvalue()

    stem = re.sub(r"\.[^.]+$", "", original_filename)
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    filename = f"{stem}_cleaned_{ts}.xlsx"

    return {
        "ok": True,
        "bytes": excel_bytes,
        "filename": filename,
        "mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }


def export_report(
    profile_before: dict,
    profile_after: dict,
    score_before: dict,
    score_after: dict,
    operation_log: list,
    original_filename: str,
) -> dict:
    """Generate an HTML cleaning report."""
    ts = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")

    ops_html = ""
    for i, entry in enumerate(operation_log, start=1):
        examples_html = ""
        for ex in entry.get("examples", []):
            examples_html += f"""
            <tr>
              <td class="ex-col">{ex.get('column','')}</td>
              <td class="before">{_esc(str(ex.get('before','')))}</td>
              <td class="after">{_esc(str(ex.get('after','')))}</td>
            </tr>"""
        ex_table = f"""
        <table class="examples-table">
          <thead><tr><th>Column</th><th>Before</th><th>After</th></tr></thead>
          <tbody>{examples_html}</tbody>
        </table>""" if examples_html else ""

        ops_html += f"""
        <div class="op-card">
          <div class="op-header">
            <span class="op-num">#{i}</span>
            <span class="op-name">{entry.get('op','').replace('_',' ').title()}</span>
            <span class="op-meta">{entry.get('rows_affected',0)} rows affected</span>
            <span class="op-ts">{entry.get('timestamp','')}</span>
          </div>
          {f'<div class="op-cols">Columns: {", ".join(entry.get("columns", []))}</div>' if entry.get("columns") else ""}
          {ex_table}
        </div>"""

    score_b = score_before.get("score", 0)
    score_a = score_after.get("score", 0)
    delta = score_a - score_b
    delta_cls = "positive" if delta >= 0 else "negative"
    delta_str = f"+{delta}" if delta >= 0 else str(delta)

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>RowBot Cleaning Report – {_esc(original_filename)}</title>
<style>
  body {{ font-family: 'Segoe UI', system-ui, sans-serif; background:#f5f7fa; color:#1a2340; margin:0; padding:24px; }}
  .container {{ max-width:900px; margin:0 auto; }}
  header {{ background:linear-gradient(135deg,#1E3A5F,#0ea5e9); color:#fff; border-radius:12px; padding:32px; margin-bottom:24px; }}
  header h1 {{ margin:0 0 8px; font-size:2rem; }}
  header p {{ margin:0; opacity:.8; }}
  .cards {{ display:grid; grid-template-columns:repeat(auto-fit,minmax(180px,1fr)); gap:16px; margin-bottom:24px; }}
  .card {{ background:#fff; border-radius:10px; padding:20px; box-shadow:0 1px 6px rgba(0,0,0,.08); }}
  .card .label {{ font-size:.8rem; color:#6b7280; text-transform:uppercase; letter-spacing:.05em; margin-bottom:6px; }}
  .card .value {{ font-size:1.8rem; font-weight:700; color:#1E3A5F; }}
  .score-delta {{ font-size:1.2rem; font-weight:600; }}
  .positive {{ color:#22c55e; }} .negative {{ color:#ef4444; }}
  .section {{ background:#fff; border-radius:10px; padding:24px; margin-bottom:20px; box-shadow:0 1px 6px rgba(0,0,0,.08); }}
  .section h2 {{ margin:0 0 16px; font-size:1.1rem; color:#1E3A5F; border-bottom:2px solid #e5e7eb; padding-bottom:8px; }}
  .op-card {{ border:1px solid #e5e7eb; border-radius:8px; padding:14px; margin-bottom:12px; }}
  .op-header {{ display:flex; align-items:center; gap:10px; flex-wrap:wrap; }}
  .op-num {{ background:#1E3A5F; color:#fff; border-radius:50%; width:24px; height:24px; display:flex; align-items:center; justify-content:center; font-size:.75rem; font-weight:700; }}
  .op-name {{ font-weight:600; flex:1; }}
  .op-meta {{ font-size:.8rem; background:#dbeafe; color:#1d4ed8; border-radius:4px; padding:2px 8px; }}
  .op-ts {{ font-size:.75rem; color:#9ca3af; }}
  .op-cols {{ font-size:.8rem; color:#6b7280; margin-top:6px; }}
  .examples-table {{ width:100%; border-collapse:collapse; margin-top:10px; font-size:.85rem; }}
  .examples-table th {{ background:#f1f5f9; text-align:left; padding:6px 10px; }}
  .examples-table td {{ padding:5px 10px; border-bottom:1px solid #f1f5f9; }}
  .before {{ color:#ef4444; font-family:monospace; }} .after {{ color:#22c55e; font-family:monospace; }}
  footer {{ text-align:center; color:#9ca3af; font-size:.8rem; margin-top:24px; }}
</style>
</head>
<body>
<div class="container">
  <header>
    <h1>🤖 RowBot Cleaning Report</h1>
    <p>File: {_esc(original_filename)} &nbsp;·&nbsp; Generated: {ts}</p>
  </header>

  <div class="cards">
    <div class="card"><div class="label">Rows (Before)</div><div class="value">{profile_before.get('row_count','-')}</div></div>
    <div class="card"><div class="label">Rows (After)</div><div class="value">{profile_after.get('row_count','-')}</div></div>
    <div class="card"><div class="label">Quality Score Before</div><div class="value">{score_b}</div></div>
    <div class="card"><div class="label">Quality Score After</div><div class="value">{score_a}</div></div>
    <div class="card"><div class="label">Score Change</div><div class="value score-delta {delta_cls}">{delta_str}</div></div>
    <div class="card"><div class="label">Operations Applied</div><div class="value">{len(operation_log)}</div></div>
  </div>

  <div class="section">
    <h2>Operations Applied</h2>
    {ops_html if ops_html else '<p style="color:#9ca3af">No operations recorded.</p>'}
  </div>

  <footer>Generated by RowBot · Messy Data In. Clean Data Out.</footer>
</div>
</body>
</html>"""

    stem = re.sub(r"\.[^.]+$", "", original_filename)
    ts_file = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    filename = f"{stem}_report_{ts_file}.html"

    return {
        "ok": True,
        "bytes": html.encode("utf-8"),
        "filename": filename,
        "mimetype": "text/html",
    }


def _esc(s: str) -> str:
    """HTML-escape a string."""
    return (
        s.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )
